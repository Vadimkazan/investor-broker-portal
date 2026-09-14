import json
import os
import base64
import io
from typing import Dict, Any, List
import psycopg2
import openpyxl

TYPE_MAP = {
    'апартаменты': 'apartments',
    'квартира в новостройке': 'new_flat',
    'квартира во вторичке': 'secondary_flat',
    'дома, дачи, коттеджи, виллы': 'houses',
    'коммерческая недвижимость': 'commercial',
    'земельные участки': 'land',
    'недвижимость за рубежом': 'abroad',
    'парковочные места': 'parking',
    'кладовые помещения': 'storage',
    'комнаты': 'rooms',
    'габы': 'gab',
    'земля под девелопмент': 'land_development',
    'здания под редевелопмент': 'buildings_redevelopment',
}

STATUS_MAP = {
    'свободен': 'available',
    'бронь': 'reserved',
    'продано': 'sold',
}

EXAMPLE_TITLE = 'апартаменты в жк «престиж»'


def escape_sql(value):
    if value is None:
        return 'NULL'
    if isinstance(value, bool):
        return 'true' if value else 'false'
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, str):
        return "'" + value.replace("'", "''").replace('\\', '\\\\') + "'"
    return "'" + str(value).replace("'", "''").replace('\\', '\\\\') + "'"


def parse_number(value) -> float:
    if value is None or value == '':
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).replace(' ', '').replace('\xa0', '').replace(',', '.')
    s = ''.join(ch for ch in s if ch.isdigit() or ch == '.')
    try:
        return float(s) if s else 0.0
    except ValueError:
        return 0.0


def parse_row(row: tuple) -> Dict[str, Any]:
    title = str(row[0] or '').strip()
    type_label = str(row[1] or '').strip().lower()
    city = str(row[2] or '').strip()
    address = str(row[3] or '').strip()
    area = parse_number(row[4])
    price = parse_number(row[5])
    yield_percent = parse_number(row[6])
    payback_years = parse_number(row[7])
    status_label = str(row[8] or '').strip().lower()
    description = str(row[9] or '').strip() if len(row) > 9 else ''
    photos_raw = str(row[10] or '').strip() if len(row) > 10 else ''

    images = [u.strip() for u in photos_raw.split(',') if u.strip()] if photos_raw else []
    property_type = TYPE_MAP.get(type_label, 'apartments')
    status = STATUS_MAP.get(status_label, 'available')

    return {
        'title': title,
        'property_type': property_type,
        'city': city,
        'address': address,
        'area': area,
        'price': price,
        'yield_percent': yield_percent,
        'payback_years': payback_years,
        'status': status,
        'description': description,
        'images': images,
    }


def handler(event: Dict[str, Any], context: Any) -> Dict[str, Any]:
    '''
    Business: Импорт объектов недвижимости брокером из заполненного Excel-файла (шаблона)
    Args: event with httpMethod, body (base64 xlsx file + broker_id)
    Returns: HTTP response with import result (created count, errors)
    '''
    method = event.get('httpMethod', 'POST')

    if method == 'OPTIONS':
        return {
            'statusCode': 200,
            'headers': {
                'Access-Control-Allow-Origin': '*',
                'Access-Control-Allow-Methods': 'POST, OPTIONS',
                'Access-Control-Allow-Headers': 'Content-Type',
                'Access-Control-Max-Age': '86400'
            },
            'body': '',
            'isBase64Encoded': False
        }

    if method != 'POST':
        return {
            'statusCode': 405,
            'headers': {'Access-Control-Allow-Origin': '*', 'Content-Type': 'application/json'},
            'body': json.dumps({'error': 'Method not allowed'})
        }

    body = json.loads(event.get('body', '{}'))
    broker_id = body.get('broker_id')
    file_b64 = body.get('file', '')

    if not broker_id:
        return {
            'statusCode': 400,
            'headers': {'Access-Control-Allow-Origin': '*', 'Content-Type': 'application/json'},
            'body': json.dumps({'error': 'broker_id required'})
        }

    if not file_b64:
        return {
            'statusCode': 400,
            'headers': {'Access-Control-Allow-Origin': '*', 'Content-Type': 'application/json'},
            'body': json.dumps({'error': 'file required'})
        }

    if ',' in file_b64:
        file_b64 = file_b64.split(',', 1)[1]

    try:
        file_bytes = base64.b64decode(file_b64)
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet_name = 'Объекты' if 'Объекты' in wb.sheetnames else wb.sheetnames[0]
        ws = wb[sheet_name]
    except Exception as e:
        return {
            'statusCode': 400,
            'headers': {'Access-Control-Allow-Origin': '*', 'Content-Type': 'application/json'},
            'body': json.dumps({'error': f'Не удалось прочитать файл: {str(e)}'})
        }

    rows_to_import: List[Dict[str, Any]] = []
    errors: List[str] = []

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        title_raw = str(row[0] or '').strip()
        if not title_raw:
            continue
        if title_raw.lower() == EXAMPLE_TITLE:
            continue

        parsed = parse_row(row)

        row_errors = []
        if not parsed['title']:
            row_errors.append('не указано название')
        if not parsed['city']:
            row_errors.append('не указан город')
        if not parsed['address']:
            row_errors.append('не указан адрес')
        if parsed['area'] <= 0:
            row_errors.append('не указана площадь')
        if parsed['price'] <= 0:
            row_errors.append('не указана цена')

        if row_errors:
            errors.append(f"Строка {idx}: {', '.join(row_errors)}")
            continue

        rows_to_import.append(parsed)

    if not rows_to_import:
        return {
            'statusCode': 400,
            'headers': {'Access-Control-Allow-Origin': '*', 'Content-Type': 'application/json'},
            'body': json.dumps({'error': 'Не найдено ни одной корректной строки для импорта', 'row_errors': errors})
        }

    conn = None
    created = 0
    try:
        dsn = os.environ.get('DATABASE_URL')
        conn = psycopg2.connect(dsn)
        conn.autocommit = True
        cur = conn.cursor()

        for item in rows_to_import:
            images_json = escape_sql(json.dumps(item['images']))
            query = f"""
                INSERT INTO investment_objects (
                    broker_id, title, city, address, property_type, area, price,
                    yield_percent, payback_years, description, images, status
                ) VALUES (
                    {escape_sql(int(broker_id))},
                    {escape_sql(item['title'])},
                    {escape_sql(item['city'])},
                    {escape_sql(item['address'])},
                    {escape_sql(item['property_type'])},
                    {escape_sql(item['area'])},
                    {escape_sql(item['price'])},
                    {escape_sql(item['yield_percent'])},
                    {escape_sql(item['payback_years'])},
                    {escape_sql(item['description'])},
                    ARRAY(SELECT json_array_elements_text({images_json}::json)),
                    {escape_sql(item['status'])}
                )
            """
            cur.execute(query)
            created += 1
    finally:
        if conn:
            conn.close()

    return {
        'statusCode': 200,
        'headers': {'Access-Control-Allow-Origin': '*', 'Content-Type': 'application/json'},
        'body': json.dumps({'created': created, 'row_errors': errors})
    }
